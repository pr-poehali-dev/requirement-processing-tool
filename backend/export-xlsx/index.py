"""Экспорт результатов анализа в XLSX с полной документацией."""
import json
import os
import base64
import io
import psycopg2
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def get_conn():
    return psycopg2.connect(os.environ["DATABASE_URL"])


STATUS_LABELS = {
    "match": "Совпадает",
    "partial": "Частично",
    "new": "Новое",
    "conflict": "Конфликт",
}

STATUS_COLORS = {
    "match": "C6EFCE",
    "partial": "FFEB9C",
    "new": "BDD7EE",
    "conflict": "FFC7CE",
}


def handler(event: dict, context) -> dict:
    headers = {
        "Access-Control-Allow-Origin": "*",
        "Access-Control-Allow-Methods": "GET, OPTIONS",
        "Access-Control-Allow-Headers": "Content-Type",
        "Content-Type": "application/json",
    }

    if event.get("httpMethod") == "OPTIONS":
        return {"statusCode": 200, "headers": headers, "body": ""}

    qs = event.get("queryStringParameters") or {}
    session_id = qs.get("session_id")

    conn = get_conn()
    cur = conn.cursor()

    if session_id:
        cur.execute("""
            SELECT ar.code, ar.requirement_text, ar.component, ar.source_file,
                   ar.status, ar.match_percent, rd.requirement as matched_text,
                   ar.analyst_comment, ar.manual_checked
            FROM analysis_results ar
            LEFT JOIN requirements_db rd ON rd.id = ar.matched_db_id
            WHERE ar.session_id = %s
            ORDER BY ar.id
        """, (session_id,))
        col_names = [d[0] for d in cur.description]
        rows = [dict(zip(col_names, r)) for r in cur.fetchall()]
        cur.execute("SELECT file_name, created_at FROM analysis_sessions WHERE id=%s", (session_id,))
        sess = cur.fetchone()
        sheet_title = f"Анализ: {sess[0]}" if sess else "Результаты анализа"
    else:
        # Экспорт всей базы требований
        cur.execute("""
            SELECT product, requirement_group, requirement, synonyms, presence,
                   check_date, risk_comment, moi_office, r7_office, is_new,
                   proposed_wording, source, external_id, author, jira_link
            FROM requirements_db
            ORDER BY product, requirement_group, id
        """)
        col_names = [d[0] for d in cur.description]
        rows = [dict(zip(col_names, r)) for r in cur.fetchall()]
        sheet_title = "База требований"

    cur.close()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(fill_type="solid", fgColor="1E3A5F")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    if session_id:
        headers_row = ["Код", "Требование", "Компонент", "Источник",
                       "Статус", "Совпадение %", "Совпало с (база)", "Комментарий аналитика", "Проверено"]
    else:
        headers_row = ["Продукт", "Группа требований", "Требование", "Синонимы",
                       "Наличие", "Дата проверки", "Комментарий-риск",
                       "Мой офис", "Р7 офис", "Новое", "Предлагаемая формулировка",
                       "Источник", "ИД", "Автор", "Ссылка на Jira"]

    for col_i, hdr in enumerate(headers_row, 1):
        cell = ws.cell(row=1, column=col_i, value=hdr)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = cell_border

    ws.row_dimensions[1].height = 30

    cell_align = Alignment(vertical="top", wrap_text=True)
    for row_i, row in enumerate(rows, 2):
        if session_id:
            vals = [
                row.get("code", ""),
                row.get("requirement_text", ""),
                row.get("component", ""),
                row.get("source_file", ""),
                STATUS_LABELS.get(row.get("status", ""), row.get("status", "")),
                row.get("match_percent", 0),
                row.get("matched_text", ""),
                row.get("analyst_comment", ""),
                "Да" if row.get("manual_checked") else "Нет",
            ]
            status_key = row.get("status", "new")
        else:
            cd = row.get("check_date")
            vals = [
                row.get("product", ""),
                row.get("requirement_group", ""),
                row.get("requirement", ""),
                row.get("synonyms", ""),
                row.get("presence", ""),
                str(cd) if cd else "",
                row.get("risk_comment", ""),
                row.get("moi_office", ""),
                row.get("r7_office", ""),
                "Да" if row.get("is_new") else "",
                row.get("proposed_wording", ""),
                row.get("source", ""),
                row.get("external_id", ""),
                row.get("author", ""),
                row.get("jira_link", ""),
            ]
            status_key = None

        for col_i, val in enumerate(vals, 1):
            cell = ws.cell(row=row_i, column=col_i, value=val)
            cell.alignment = cell_align
            cell.border = cell_border
            if status_key and col_i == 5:
                cell.fill = PatternFill(fill_type="solid", fgColor=STATUS_COLORS.get(status_key, "FFFFFF"))

    # Ширина колонок
    col_widths = [12, 55, 25, 20, 15, 12, 45, 35, 10] if session_id else \
                 [20, 25, 50, 25, 15, 14, 30, 20, 20, 8, 35, 20, 15, 20, 30]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    xlsx_b64 = base64.b64encode(buf.getvalue()).decode()

    return {
        "statusCode": 200,
        "headers": {**headers, "Content-Type": "application/json"},
        "body": json.dumps({
            "file_base64": xlsx_b64,
            "file_name": f"{'analysis_' + str(session_id) if session_id else 'requirements_db'}.xlsx",
        })
    }
